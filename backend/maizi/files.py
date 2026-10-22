from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import defusedxml  # noqa: F401 -- enables openpyxl's protected XML integration
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException

from maizi.numbers import positive_decimal

PARSER_VERSION = 2
MAX_CELL_CHARS = 32767
MAX_BYTES = 5 * 1024 * 1024
MAX_COLUMNS = 100
MAX_PHYSICAL_ROWS = 20001


def _safe_filename(filename) -> str:
    text = str(filename).replace("\\", "/")
    name = text.rsplit("/", 1)[-1]
    return name or "input"


def _input_error(message: str, code: str, line_number: int) -> ValueError:
    error = ValueError(message)
    error.code = code
    error.line_number = line_number
    return error


def _cell_error(filename, line_number: int) -> ValueError:
    safe = _safe_filename(filename)
    return _input_error(
        f"{safe}: physical row {line_number} contains a cell longer than "
        f"{MAX_CELL_CHARS} characters",
        "CELL_TOO_LONG",
        line_number,
    )


def _raw_error(filename, line_number: int) -> ValueError:
    safe = _safe_filename(filename)
    return _input_error(
        f"{safe}: physical row {line_number} has serialized raw data beyond "
        f"the Excel cell limit of {MAX_CELL_CHARS} characters",
        "RAW_TOO_LONG",
        line_number,
    )


def _csv_quote_error(filename, line_number: int) -> ValueError:
    safe = _safe_filename(filename)
    return _input_error(
        f"{safe}: malformed CSV quoting at physical row {line_number}",
        "INVALID_CSV",
        line_number,
    )


def _check_cell_lengths(values, filename, line_number: int) -> None:
    for value in values:
        if len(str(value)) > MAX_CELL_CHARS:
            raise _cell_error(filename, line_number)


def _check_raw_length(data: dict, filename, line_number: int) -> None:
    if len(json.dumps(data, ensure_ascii=False)) > MAX_CELL_CHARS:
        raise _raw_error(filename, line_number)


def _validate_csv_quotes(text: str, dialect, filename) -> None:
    """Reject quotes that are not valid RFC-style CSV quoting.

    The standard csv module is deliberately permissive about a bare quote in an
    unquoted field. This lightweight scanner adds that check while allowing
    embedded delimiters, embedded newlines and doubled quote escapes.
    """
    quotechar = getattr(dialect, "quotechar", None)
    if not quotechar:
        return
    delimiter = getattr(dialect, "delimiter", None) or ","
    doublequote = bool(getattr(dialect, "doublequote", True))
    if getattr(dialect, "escapechar", None):
        # Let the csv module handle explicit escape characters.
        return

    line = 1
    index = 0
    length = len(text)
    in_quotes = False
    closed_quote = False
    field_start = True

    while index < length:
        char = text[index]
        if in_quotes:
            if char == quotechar:
                if doublequote and index + 1 < length and text[index + 1] == quotechar:
                    index += 2
                    continue
                in_quotes = False
                closed_quote = True
                index += 1
                continue
            if char == "\r":
                if index + 1 < length and text[index + 1] == "\n":
                    index += 1
                line += 1
            elif char == "\n":
                line += 1
            index += 1
            continue

        if closed_quote:
            if char == delimiter:
                closed_quote = False
                field_start = True
                index += 1
                continue
            if char == "\r":
                if index + 1 < length and text[index + 1] == "\n":
                    index += 1
                line += 1
                closed_quote = False
                field_start = True
                index += 1
                continue
            if char == "\n":
                line += 1
                closed_quote = False
                field_start = True
                index += 1
                continue
            raise _csv_quote_error(filename, line)

        if char == quotechar:
            if field_start:
                in_quotes = True
                field_start = False
                index += 1
                continue
            raise _csv_quote_error(filename, line)
        if char == delimiter:
            field_start = True
            index += 1
            continue
        if char == "\r":
            if index + 1 < length and text[index + 1] == "\n":
                index += 1
            line += 1
            field_start = True
            index += 1
            continue
        if char == "\n":
            line += 1
            field_start = True
            index += 1
            continue
        field_start = False
        index += 1

    if in_quotes:
        raise _csv_quote_error(filename, line)


def value_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        fmt = cell.number_format or ""
        if re.fullmatch(r"0{2,}", fmt) and float(value).is_integer():
            return str(int(value)).zfill(len(fmt))
    return str(value)


def parse_table(content: bytes, filename: str, role: str, sheet_name=None):
    if len(content) > MAX_BYTES:
        raise ValueError("Each file must be at most 5 MB")
    extension = Path(filename).suffix.lower()
    physical, sheets = [], []
    if extension == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV must be UTF-8 or UTF-8 BOM")
        if "\x00" in text:
            raise ValueError("Invalid CSV content")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        _validate_csv_quotes(text, dialect, filename)
        try:
            csv.field_size_limit(MAX_BYTES)
        except OverflowError:
            pass
        reader = csv.reader(io.StringIO(text), dialect, strict=True)
        previous_end = 0
        try:
            for values in reader:
                start = previous_end + 1
                end = max(start, reader.line_num)
                previous_end = end
                if end > MAX_PHYSICAL_ROWS:
                    raise ValueError("Table exceeds the physical-row safety limit")
                if len(values) > MAX_COLUMNS:
                    raise ValueError("At most 100 columns are supported")
                _check_cell_lengths(values, filename, start)
                physical.append((start, end, [str(x) for x in values], []))
        except csv.Error as exc:
            line_number = max(1, reader.line_num or previous_end + 1)
            raise _csv_quote_error(filename, line_number) from exc
        selected = "CSV"
    elif extension == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Workbook exceeds the decompression safety limit")
                if any("vbaProject" in x.filename for x in archive.infolist()):
                    raise ValueError("Macros are not supported")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheets = workbook.sheetnames
                if len(sheets) > 10:
                    raise ValueError("At most ten worksheets are supported")
                if sheet_name and sheet_name not in sheets:
                    raise ValueError("Selected worksheet does not exist")
                if len(sheets) > 1 and not sheet_name:
                    raise ValueError("Choose one worksheet explicitly: " + ", ".join(sheets))
                sheet = workbook[sheet_name or sheets[0]]
                selected = sheet.title
                if (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ValueError("At most 100 columns are supported")
                for number, cells in enumerate(sheet.iter_rows(), 1):
                    if number > MAX_PHYSICAL_ROWS:
                        raise ValueError("Table exceeds the physical-row safety limit")
                    values = [value_text(cell) for cell in cells]
                    formulas = [index for index, cell in enumerate(cells) if cell.data_type == "f"]
                    _check_cell_lengths(values, filename, number)
                    physical.append((number, number, values, formulas))
            finally:
                workbook.close()
        except (zipfile.BadZipFile, KeyError, OSError, ParseError, DefusedXmlException, InvalidFileException) as exc:
            raise ValueError("Workbook cannot be read") from exc
    else:
        raise ValueError("Only CSV and XLSX files are supported")

    if not physical or not any(physical[0][2]):
        raise ValueError("The first row must contain the column headers")
    headers = [x.strip() for x in physical[0][2]]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not x for x in headers) or len(set(headers)) != len(headers):
        raise ValueError("Headers must be non-empty and unique")
    if physical[0][3]:
        raise ValueError("Formula headers are not supported")
    records, blanks = [], 0
    for start, end, values, formulas in physical[1:]:
        if not any(str(x).strip() for x in values):
            blanks += 1
            continue
        if len(values) > len(headers) and any(str(x).strip() for x in values[len(headers):]):
            raise ValueError("A data row contains values outside the header columns")
        if formulas and role in {"master", "template"}:
            raise ValueError("Master and template formulas are not supported")
        data = {
            key: values[i] if i < len(values) else ""
            for i, key in enumerate(headers)
        }
        _check_raw_length(data, filename, start)
        records.append(
            {
                "row_number": start,
                "row_end_number": end,
                "data": data,
                "formula_fields": [headers[i] for i in formulas if i < len(headers)],
            }
        )
    limit = 5000 if role == "master" else 10000
    if len(records) > limit:
        raise ValueError("This table exceeds the allowed row count")
    if role != "template" and not records:
        raise ValueError("The file has headers but no data rows")
    return {
        "headers": headers,
        "records": records,
        "sheet": selected,
        "sheets": sheets or ["CSV"],
        "blank_rows": blanks,
        "parser_version": PARSER_VERSION,
    }


def profile(table):
    result = {}
    for key in table["headers"]:
        values, missing = {}, 0
        for row in table["records"]:
            value = str(row["data"].get(key, "")).strip()
            if not value:
                missing += 1
            else:
                values[value] = values.get(value, 0) + 1
        result[key] = {
            "missing": missing,
            "distinct_count": len(values),
            "values": dict(list(values.items())[:60]),
        }
    return {
        "headers": table["headers"],
        "row_count": len(table["records"]),
        "columns": result,
        "samples": [r["data"] for r in table["records"][:5]],
        "sheet": table["sheet"],
        "blank_rows": table["blank_rows"],
    }


def extract_products(table, mapping):
    required = ["sku", "name", "base_unit"]
    if any(not mapping.get(x) or mapping[x] not in table["headers"] for x in required):
        raise ValueError("Confirm the SKU, name and base-unit fields in the product master")
    products, seen = [], set()
    for row in table["records"]:
        data = row["data"]
        sku = str(data.get(mapping["sku"], "")).strip()
        if not sku or sku in seen:
            raise ValueError("Master SKU must be present and unique")
        seen.add(sku)
        base = str(data.get(mapping["base_unit"], "")).strip()
        if not base:
            raise ValueError("Master base unit is required")
        factors = {base: "1"}
        for unit, column in mapping.items():
            if unit in required or not column or column not in table["headers"]:
                continue
            raw = str(data.get(column, "")).strip()
            if not raw:
                continue
            try:
                factor = positive_decimal(raw, Decimal("1000000000"))
                if not factor.is_finite() or factor <= 0:
                    raise ValueError()
            except (InvalidOperation, ValueError):
                raise ValueError("Master conversion must be positive, at most 1e9, with at most six decimal places")
            factors[unit] = format(factor, "f")
        products.append(
            {
                "sku": sku,
                "name": str(data.get(mapping["name"], "")).strip(),
                "base_unit": base,
                "unit_factors": factors,
                "master_row_number": row["row_number"],
            }
        )
    return products
