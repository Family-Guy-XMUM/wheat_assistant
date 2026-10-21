import json
from copy import copy
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

def literal(sheet, row):
    if any(isinstance(value, str) and len(value) > 32767 for value in row):
        raise ValueError("An output cell exceeds the Excel text limit. Narrow the input; no truncated workbook was delivered.")
    # Excel numeric cells have limited precision. Keep large exact integers as text.
    row = [str(value) if isinstance(value, int) and len(str(abs(value))) > 15 else value for value in row]
    sheet.append(row)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
            cell.number_format = "@"

def finish(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="343A35")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(65, max(14, max(len(str(c.value or "")) for c in list(column)[:100]) + 2))

def export_workbook(result, flow, snapshot, target, template_content=None):
    workbook = Workbook()
    pick = workbook.active
    pick.title = "Warehouse_Picklist"
    columns = flow["output_columns"]
    literal(pick, [x["header"] for x in columns])
    if template_content:
        template = load_workbook(BytesIO(template_content), read_only=False, data_only=False)
        source = template[template.sheetnames[0]]
        for index, cell in enumerate(source[1], 1):
            if index <= len(columns):
                pick.cell(1, index)._style = copy(cell._style)
        template.close()
    for item in result["sku_summary"]:
        literal(pick, [int(item[x["field"]]) if x["field"] == "quantity" else item[x["field"]] for x in columns])
    detail_headers = ["Source row", "Channel", "File", "Sheet", "Row", "SKU", "Product", "Base quantity", "Unit", "Disposition", "Reasons", "Duplicate of"]
    def record(row):
        return [row["source_row_id"], row["source_id"], row.get("source_name", ""), row.get("sheet", ""),
                row.get("row_number"), row.get("sku"), row.get("product_name", ""), row.get("quantity"),
                row.get("unit", ""), row["disposition"], "; ".join(x["code"] + ": " + x["message"] for x in row["reasons"]),
                row.get("duplicate_of")]
    for name, disposition in [("Order_Detail", "included"), ("Exceptions", "exception"), ("Excluded", "excluded"), ("Duplicates", "duplicate")]:
        sheet = workbook.create_sheet(name)
        literal(sheet, detail_headers + (["Original values"] if disposition == "exception" else []))
        for row in result["rows"]:
            if row["disposition"] == disposition:
                literal(sheet, record(row) + ([json.dumps(row["raw"], ensure_ascii=False)] if disposition == "exception" else []))
    summary = workbook.create_sheet("SKU_Summary")
    literal(summary, ["SKU", "Name", "Quantity", "Unit"])
    for item in result["sku_summary"]:
        literal(summary, [item["sku"], item["name"], int(item["quantity"]), item["unit"]])
    lineage = workbook.create_sheet("Lineage")
    literal(lineage, ["Source row", "Disposition", "Original values", "Processing events"])
    for row in result["rows"]:
        literal(lineage, [row["source_row_id"], row["disposition"], json.dumps(row["raw"], ensure_ascii=False),
                          json.dumps(row["events"], ensure_ascii=False)])
    audit = workbook.create_sheet("Run_Summary")
    literal(audit, ["Item", "Value"])
    for key, value in {**result["summary"], "snapshot": snapshot, "scope": "Picking preparation only", "orbyte": "Simulation"}.items():
        literal(audit, [key, str(value)])
    for sheet in workbook:
        finish(sheet)
    workbook.properties.creator = "MaiZi"
    workbook.properties.lastModifiedBy = "MaiZi"
    Path(target).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return str(target)
