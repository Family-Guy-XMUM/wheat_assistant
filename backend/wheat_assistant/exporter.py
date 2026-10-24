from __future__ import annotations

import json
import os
import tempfile
from copy import copy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from wheat_assistant.files import MAX_CELL_CHARS

EXPORTER_VERSION = 2
SHEET_NAMES = [
    "Warehouse_Picklist",
    "Order_Detail",
    "Exceptions",
    "Excluded",
    "Duplicates",
    "SKU_Summary",
    "Lineage",
    "Run_Summary",
]


def literal(sheet, row, row_number=None):
    row = list(row)
    if any(isinstance(value, str) and len(value) > MAX_CELL_CHARS for value in row):
        raise ValueError(
            "An output cell exceeds the Excel text limit. "
            "Narrow the input; no truncated workbook was delivered."
        )
    # Excel numeric cells have limited precision. Keep large exact integers as text.
    row = [
        str(value)
        if isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15
        else value
        for value in row
    ]
    if row_number is None:
        row_number = sheet.max_row + 1
    for column, value in enumerate(row, 1):
        cell = sheet.cell(row=row_number, column=column, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
            cell.number_format = "@"
    return row_number


def finish(
    sheet,
    *,
    default_header=True,
    auto_width=True,
    preserve_widths=frozenset(),
    header_height=24,
):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if default_header:
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="343A35")
            cell.alignment = Alignment(vertical="center")
    if header_height is not None:
        sheet.row_dimensions[1].height = header_height
    if auto_width:
        for column in sheet.columns:
            letter = column[0].column_letter
            if letter in preserve_widths:
                continue
            sheet.column_dimensions[letter].width = min(
                65,
                max(
                    14,
                    max(len(str(c.value or "")) for c in list(column)[:100]) + 2,
                ),
            )


def _copy_public_style(source, target):
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    target.number_format = source.number_format


def _apply_template_header(sheet, source, column_count):
    copied_widths = set()
    for index in range(1, min(column_count, source.max_column or 0) + 1):
        _copy_public_style(source.cell(row=1, column=index), sheet.cell(row=1, column=index))
    for letter, dimension in source.column_dimensions.items():
        target_dimension = sheet.column_dimensions[letter]
        if dimension.width is not None:
            target_dimension.width = dimension.width
            copied_widths.add(letter)
        target_dimension.hidden = dimension.hidden
        target_dimension.bestFit = dimension.bestFit
    height = source.row_dimensions[1].height
    if height is not None:
        sheet.row_dimensions[1].height = height
    return frozenset(copied_widths)


def _expected_row_counts(result, snapshot):
    rows = result["rows"]
    summary_count = len(
        {
            **result["summary"],
            "snapshot": snapshot,
            "scope": "Picking preparation only",
            "orbyte": "Simulation",
        }
    )
    return {
        "Warehouse_Picklist": 1 + len(result["sku_summary"]),
        "Order_Detail": 1 + sum(1 for row in rows if row["disposition"] == "included"),
        "Exceptions": 1 + sum(1 for row in rows if row["disposition"] == "exception"),
        "Excluded": 1 + sum(1 for row in rows if row["disposition"] == "excluded"),
        "Duplicates": 1 + sum(1 for row in rows if row["disposition"] == "duplicate"),
        "SKU_Summary": 1 + len(result["sku_summary"]),
        "Lineage": 1 + len(rows),
        "Run_Summary": 1 + summary_count,
    }


def _verify_export(path, expected_rows):
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("Exported workbook verification failed") from exc
    try:
        if workbook.sheetnames != SHEET_NAMES:
            raise ValueError("Exported workbook has unexpected sheets")
        for name in SHEET_NAMES:
            sheet = workbook[name]
            if sheet.max_row != expected_rows[name]:
                raise ValueError("Exported workbook has unexpected row counts")
            header = next(sheet.iter_rows(min_row=1, max_row=1), ())
            for cell in header:
                try:
                    _ = (
                        cell.font,
                        cell.fill,
                        cell.border,
                        cell.alignment,
                        cell.protection,
                        cell.number_format,
                    )
                except Exception as exc:
                    raise ValueError("Exported workbook contains an invalid style") from exc
        # Stream every worksheet XML member instead of constructing another
        # ReadOnlyCell for every output cell merely to check its formula type.
        # This also covers cells beyond an incorrect worksheet dimension.
        with ZipFile(path) as archive:
            worksheets = [info for info in archive.infolist()
                          if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml")]
            if len(worksheets) != len(SHEET_NAMES) or len({info.filename for info in worksheets}) != len(SHEET_NAMES):
                raise ValueError("Exported workbook has unexpected worksheet parts")
            for info in worksheets:
                with archive.open(info) as stream:
                    for _, element in iterparse(stream, events=("end",), forbid_dtd=True):
                        if element.tag.rsplit("}", 1)[-1] == "f":
                            raise ValueError("Exported workbook contains formulas")
                        element.clear()
    finally:
        workbook.close()


def export_workbook(
    result,
    flow,
    snapshot,
    target,
    template_content=None,
    template_sheet_name=None,
    on_stage=None,
):
    target_path = Path(target)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    pick = workbook.active
    pick.title = "Warehouse_Picklist"
    columns = flow["output_columns"]
    literal(pick, [column["header"] for column in columns], 1)

    template_header = False
    preserved_widths = frozenset()
    if template_content:
        template = load_workbook(BytesIO(template_content), read_only=False, data_only=False)
        try:
            if template_sheet_name:
                if template_sheet_name not in template.sheetnames:
                    raise ValueError("Selected template worksheet does not exist")
                source = template[template_sheet_name]
            else:
                source = template[template.sheetnames[0]]
            preserved_widths = _apply_template_header(pick, source, len(columns))
            template_header = True
        finally:
            template.close()

    row_index = 2
    for item in result["sku_summary"]:
        literal(
            pick,
            [
                int(item[column["field"]])
                if column["field"] == "quantity"
                else item[column["field"]]
                for column in columns
            ],
            row_index,
        )
        row_index += 1
    finish(
        pick,
        default_header=not template_header,
        auto_width=True,
        preserve_widths=preserved_widths,
        header_height=None if template_header else 24,
    )

    detail_headers = [
        "Source row",
        "Channel",
        "File",
        "Sheet",
        "Row",
        "SKU",
        "Product",
        "Base quantity",
        "Unit",
        "Disposition",
        "Reasons",
        "Duplicate of",
    ]

    def record(row):
        return [
            row["source_row_id"],
            row["source_id"],
            row.get("source_name", ""),
            row.get("sheet", ""),
            row.get("row_number"),
            row.get("sku"),
            row.get("product_name", ""),
            row.get("quantity"),
            row.get("unit", ""),
            row["disposition"],
            "; ".join(reason["code"] + ": " + reason["message"] for reason in row["reasons"]),
            row.get("duplicate_of"),
        ]

    for name, disposition in [
        ("Order_Detail", "included"),
        ("Exceptions", "exception"),
        ("Excluded", "excluded"),
        ("Duplicates", "duplicate"),
    ]:
        sheet = workbook.create_sheet(name)
        literal(
            sheet,
            detail_headers + (["Original values"] if disposition == "exception" else []),
            1,
        )
        row_index = 2
        for row in result["rows"]:
            if row["disposition"] == disposition:
                literal(
                    sheet,
                    record(row)
                    + ([json.dumps(row["raw"], ensure_ascii=False)] if disposition == "exception" else []),
                    row_index,
                )
                row_index += 1
        finish(sheet)

    summary = workbook.create_sheet("SKU_Summary")
    literal(summary, ["SKU", "Name", "Quantity", "Unit"], 1)
    row_index = 2
    for item in result["sku_summary"]:
        literal(summary, [item["sku"], item["name"], int(item["quantity"]), item["unit"]], row_index)
        row_index += 1
    finish(summary)

    lineage = workbook.create_sheet("Lineage")
    literal(
        lineage,
        ["Source row", "Disposition", "Original values", "Processing events"],
        1,
    )
    row_index = 2
    for row in result["rows"]:
        literal(
            lineage,
            [
                row["source_row_id"],
                row["disposition"],
                json.dumps(row["raw"], ensure_ascii=False),
                json.dumps(row["events"], ensure_ascii=False),
            ],
            row_index,
        )
        row_index += 1
    finish(lineage)

    audit = workbook.create_sheet("Run_Summary")
    literal(audit, ["Item", "Value"], 1)
    row_index = 2
    for key, value in {
        **result["summary"],
        "snapshot": snapshot,
        "scope": "Picking preparation only",
        "orbyte": "Simulation",
    }.items():
        literal(audit, [key, str(value)], row_index)
        row_index += 1
    finish(audit)

    workbook.properties.creator = "wheat_assistant"
    workbook.properties.lastModifiedBy = "wheat_assistant"
    expected_rows = _expected_row_counts(result, snapshot)

    temporary = None
    try:
        if on_stage is not None:
            on_stage("writing")
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{target_path.name}.",
            suffix=".xlsx",
            dir=str(target_path.parent),
        )
        os.close(descriptor)
        temporary = Path(temporary_name)
        workbook.save(temporary)
        if on_stage is not None:
            on_stage("verifying")
        _verify_export(temporary, expected_rows)
        os.replace(temporary, target_path)
    finally:
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass

    return str(target)
