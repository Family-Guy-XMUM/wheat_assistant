"""Foundation audit: durable approvals and portable Unicode worker boundaries."""
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

def ascii_defaults(monkeypatch):
    """Emulate a non-UTF-8 computer at the actual JSON file boundary."""
    original_read,original_write=Path.read_text,Path.write_text
    def read(path,encoding=None,errors=None):
        return original_read(path,encoding=encoding or "ascii",errors=errors)
    def write(path,data,encoding=None,errors=None,newline=None):
        return original_write(path,data,encoding=encoding or "ascii",errors=errors,newline=newline)
    monkeypatch.setattr(Path,"read_text",read)
    monkeypatch.setattr(Path,"write_text",write)

def unicode_case():
    rows=[{"source_id":"手工","source_row_id":"手工:2","source_name":"合成订单.csv","row_number":2,"data":{"订单":"虚构-001","明细":"1","商品":"茶-001","数量":"2","单位":"箱","状态":"已确认"}}]
    products=[{"sku":"茶-001","name":"合成茉莉茶","base_unit":"件","unit_factors":{"箱":"12"}}]
    flow={"channels":{"手工":{"fields":{"order_id":"订单","line_id":"明细","product":"商品","quantity":"数量","unit":"单位","status":"状态"},"status_actions":{"已确认":"include"}}}}
    return rows,products,flow

def assert_unicode_result(result):
    assert result["summary"]["included"]==1
    assert result["rows"][0]["raw"]["订单"]=="虚构-001"
    assert result["rows"][0]["source_name"]=="合成订单.csv"
    assert result["sku_summary"]==[{"sku":"茶-001","name":"合成茉莉茶","quantity":"24","unit":"件"}]

def test_unicode_in_real_child_with_ascii_parent_defaults(monkeypatch):
    from wheat_assistant.execution import execute_isolated
    ascii_defaults(monkeypatch)
    assert_unicode_result(execute_isolated(*unicode_case()))

def test_worker_reads_and_writes_utf8_with_ascii_defaults(tmp_path,monkeypatch):
    from wheat_assistant.worker import main
    rows,products,flow=unicode_case()
    source,target=tmp_path/"input.json",tmp_path/"result.json"
    source.write_bytes(json.dumps({"rows":rows,"products":products,"flow":flow},ensure_ascii=False).encode("utf-8"))
    ascii_defaults(monkeypatch)
    main([str(source),str(target)])
    assert_unicode_result(json.loads(target.read_bytes().decode("utf-8")))

@pytest.fixture()
def factory(monkeypatch,tmp_path):
    import dotenv
    monkeypatch.setattr(dotenv,"load_dotenv",lambda *a,**kw:False)
    for key in ["LLM_BASE_URL","LLM_MODEL","LLM_API_KEY"]:monkeypatch.setenv(key,"")
    from wheat_assistant.api import create_app
    return lambda:create_app(data_dir=tmp_path)

def approved_run(client):
    batch=client.post("/api/demo/load",json={"week":1}).json()
    spec=client.post(f"/api/batches/{batch['id']}/plan",json={"mode":"suggested"}).json()["flow"]
    spec["aliases"]={"Jasmine500":"TEA-001"};spec["unit_aliases"]={"carton":"box"}
    pre=client.post(f"/api/batches/{batch['id']}/preflight",json=spec).json()
    approval=client.post(f"/api/batches/{batch['id']}/approve",json={"flow":spec,"reviewed":["fields","statuses","products","units","duplicates","output"],"expected_input_hash":pre["input_hash"]})
    assert approval.status_code==200,approval.text
    saved=approval.json()
    run=client.post("/api/runs",json={"batch_id":batch["id"],"flow_id":saved["id"]}).json()
    return batch,saved,run

def test_interrupted_run_resumes_and_approved_export_survives_restarts(factory):
    first=factory()
    with TestClient(first) as client:
        batch,saved,run=approved_run(client)
        first.state.store.put("run",run["id"],{**run,"status":"running"})
    second=factory()
    with TestClient(second) as client:
        interrupted=client.get(f"/api/runs/{run['id']}").json()
        assert interrupted["status"]=="interrupted"
        response=client.post(f"/api/runs/{run['id']}/execute")
        assert response.status_code==200,response.text
        completed=response.json()
        assert completed["summary"]=={"input":36,"included":24,"excluded":3,"duplicate":3,"exception":6,"totals_match":True}
        approved=client.post(f"/api/runs/{run['id']}/export-approval",json={"result_hash":completed["result_hash"],"acknowledge_exceptions":True})
        assert approved.status_code==200,approved.text
        artifact=client.get(approved.json()["download_url"]).content
        handoff=client.post(f"/api/runs/{run['id']}/handoff").json()
    with TestClient(factory()) as client:
        restored=client.get(f"/api/runs/{run['id']}").json()
        assert restored["status"]=="completed" and restored["export_approved"] is True
        assert restored["result_hash"]==completed["result_hash"]
        same=client.post("/api/runs",json={"batch_id":batch["id"],"flow_id":saved["id"]}).json()
        assert same["id"]==run["id"]
        assert client.get(f"/api/runs/{run['id']}/download").content==artifact
        assert client.post(f"/api/runs/{run['id']}/handoff").json()["id"]==handoff["id"]
        rows=client.get(f"/api/runs/{run['id']}/rows").json()["rows"]
        assert len(rows)==36 and all(row["source_row_id"] for row in rows)
        workbook=load_workbook(io.BytesIO(artifact))
        assert workbook["Exceptions"].max_row==7
        assert workbook["Lineage"].max_row==37

def test_changed_original_bytes_invalidate_existing_approval(factory):
    app=factory()
    with TestClient(app) as client:
        batch,saved,run=approved_run(client)
        asset=app.state.store.get("file",batch["file_ids"][0])
        original=app.state.store.root/"files"/asset["storage_name"]
        original.chmod(0o600)
        original.write_bytes(original.read_bytes()+b"\n")
        response=client.post(f"/api/runs/{run['id']}/execute")
        assert response.status_code==422
        assert "original input file changed" in response.json()["detail"]
        assert client.get(f"/api/runs/{run['id']}").json()["status"]=="ready"
        assert client.get(f"/api/runs/{run['id']}/download").status_code==403
