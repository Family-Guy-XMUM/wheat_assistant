"""Generated workbook validation must inspect every worksheet, including its tail."""
import io
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from wheat_assistant.exporter import SHEET_NAMES, _verify_export, literal

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def synthetic_book():
    book = Workbook()
    book.remove(book.active)
    for name in SHEET_NAMES:
        literal(book.create_sheet(name), ["Synthetic header"], 1)
    return book


@pytest.mark.parametrize("sheet_name", ["Order_Detail", "Lineage"])
def test_verifier_rejects_a_formula_beyond_declared_dimensions(tmp_path, sheet_name):
    book = synthetic_book()
    book[sheet_name].cell(row=1002, column=1, value="=1+1")
    original = io.BytesIO()
    book.save(original)
    book.close()
    target = tmp_path / "synthetic-hidden-formula.xlsx"
    member = "xl/worksheets/sheet" + str(SHEET_NAMES.index(sheet_name) + 1) + ".xml"
    with ZipFile(io.BytesIO(original.getvalue())) as source, ZipFile(target, "w") as output:
        for info in source.infolist():
            content = source.read(info)
            if info.filename == member:
                root = ElementTree.fromstring(content)
                root.find(NS + "dimension").set("ref", "A1:A1")
                content = ElementTree.tostring(root, encoding="utf-8")
            output.writestr(info, content)
    with pytest.raises(ValueError, match="formulas"):
        _verify_export(target, {name: 1 for name in SHEET_NAMES})


def test_verifier_keeps_formula_looking_literal_text(tmp_path):
    book = synthetic_book()
    for name in SHEET_NAMES:
        literal(book[name], ["=SUM(A1:A2)"], 1002)
    target = tmp_path / "synthetic-literal-text.xlsx"
    book.save(target)
    book.close()
    _verify_export(target, {name: 1002 for name in SHEET_NAMES})
