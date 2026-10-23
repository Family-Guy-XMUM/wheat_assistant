"""Coordinator-authored boundary checks, independent of live credentials."""
import copy
import io
import json
import zipfile
from decimal import Decimal

import pytest

from wheat_assistant.engine import process_orders
from wheat_assistant.execution import execute_isolated
from wheat_assistant.files import parse_table
from wheat_assistant.numbers import positive_decimal

def setup_rows(count=1, products_count=1):
    products=[{"sku":f"SKU-{i:04d}","name":f"Synthetic product {i}","base_unit":"piece","unit_factors":{"box":"12"}} for i in range(products_count)]
    flow={"channels":{"a":{"fields":{"order_id":"Order","line_id":"Line","product":"SKU","quantity":"Qty","unit":"Unit","status":"Status"},"status_actions":{"Confirmed":"include"},"dedup_mode":"business_key"}}}
    rows=[{"source_id":"a","source_row_id":f"a:{i}","row_number":i+2,"data":{"Order":f"O-{i}","Line":"1","SKU":products[i%products_count]["sku"],"Qty":"2","Unit":"box","Status":"Confirmed"}} for i in range(count)]
    return rows,products,flow

def test_maximum_row_and_master_scale_in_real_worker():
    rows,products,flow=setup_rows(10000,5000)
    result=execute_isolated(rows,products,flow)
    assert result["summary"]=={"input":10000,"included":10000,"excluded":0,"duplicate":0,"exception":0,"totals_match":True}
    assert len(result["sku_summary"])==5000
    assert all(x["quantity"]=="48" for x in result["sku_summary"])
    assert len({x["source_row_id"] for x in result["rows"]})==10000

@pytest.mark.parametrize("value", ["1e999999", "1e-999999", "NaN", "Infinity", "1000000000001", "0.0000001"])
def test_hostile_numeric_values_are_quarantined(value):
    rows,products,flow=setup_rows()
    rows[0]["data"]["Qty"]=value
    result=process_orders(rows,products,flow)
    assert result["summary"]["exception"]==1
    assert result["sku_summary"]==[]
    assert any(x["code"]=="INVALID_QUANTITY" for x in result["rows"][0]["reasons"])

def test_exact_large_integer_aggregation():
    rows,products,flow=setup_rows(10000)
    products[0]["unit_factors"]["box"]="1000000000"
    for row in rows: row["data"]["Qty"]="1000000000000"
    result=process_orders(rows,products,flow)
    assert result["sku_summary"][0]["quantity"]=="10000000000000000000000000"

def test_malformed_workbook_xml_is_an_input_error():
    from openpyxl import Workbook
    buffer=io.BytesIO();Workbook().save(buffer)
    output=io.BytesIO()
    with zipfile.ZipFile(buffer) as source,zipfile.ZipFile(output,"w") as target:
        for item in source.infolist():
            target.writestr(item.filename,b"<broken" if item.filename=="xl/workbook.xml" else source.read(item))
    with pytest.raises(ValueError, match="Workbook cannot be read"):
        parse_table(output.getvalue(),"bad.xlsx","orders")

def test_provider_invalid_schema_repairs_once_without_executing(tmp_path, monkeypatch):
    import wheat_assistant.planner as planner
    from wheat_assistant.models import FlowSpec
    from wheat_assistant.store import Store
    for name,value in {"LLM_BASE_URL":"https://invalid.example/v1","LLM_MODEL":"test-model","LLM_API_KEY":"synthetic-test-token"}.items():monkeypatch.setenv(name,value)
    calls=[]
    class Response:
        status_code=200
        def __init__(self,content):self.content=content
        def json(self):return {"choices":[{"message":{"content":json.dumps(self.content)}}]}
    class Client:
        def __init__(self,*args,**kwargs):pass
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def post(self,*args,**kwargs):
            calls.append(kwargs["json"])
            return Response({"flow":{},"questions":"bad-type"} if len(calls)==1 else {"flow":spec,"questions":["Review this rule"]})
    monkeypatch.setattr(planner.httpx,"Client",Client)
    rows,products,raw=setup_rows()
    spec=FlowSpec.model_validate(raw).model_dump()
    assets=[{"role":"orders","channel_id":"a","profile":{},"table":{"headers":list(rows[0]["data"])}}]
    store=Store(tmp_path)
    value=planner.propose_live(store,{"id":"batch"},assets,"Synthetic test",spec)
    assert value["origin"]=="live_ai" and value["questions"]==["Review this rule"]
    assert len(calls)==2
    assert sorted(x["status"] for x in store.list("usage"))==["invalid_output","succeeded"]

def test_provider_timeout_is_visible_and_not_retried(tmp_path,monkeypatch):
    import httpx
    import wheat_assistant.planner as planner
    from wheat_assistant.store import Store
    for name in ["LLM_BASE_URL","LLM_MODEL","LLM_API_KEY"]:monkeypatch.setenv(name,"synthetic")
    class Client:
        def __init__(self,*args,**kwargs):pass
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def post(self,*args,**kwargs):raise httpx.ReadTimeout("synthetic timeout")
    monkeypatch.setattr(planner.httpx,"Client",Client)
    store=Store(tmp_path)
    with pytest.raises(ValueError,match="timed out"):
        planner.propose_live(store,{"id":"batch"},[],"",{})
    assert [x["status"] for x in store.list("usage")]==["timeout"]

def test_export_preserves_large_integers_and_refuses_truncation(tmp_path):
    from wheat_assistant.exporter import export_workbook, literal
    from wheat_assistant.models import FlowSpec
    from openpyxl import load_workbook, Workbook
    rows,products,flow=setup_rows()
    result=process_orders(rows,products,flow)
    result["sku_summary"][0]["quantity"]="1234567890123456789"
    target=tmp_path/"large.xlsx"
    export_workbook(result,FlowSpec.model_validate(flow).model_dump(),"synthetic",target)
    book=load_workbook(target)
    assert book["Warehouse_Picklist"]["C2"].value=="1234567890123456789"
    assert book["Warehouse_Picklist"]["C2"].data_type=="s"
    with pytest.raises(ValueError,match="Excel text limit"):
        literal(Workbook().active,["x"*32768])
