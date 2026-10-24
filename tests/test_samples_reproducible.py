"""The same synthetic week must remain identical across a clock rollover."""
import io
from datetime import datetime, timezone
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.writer import excel

from wheat_assistant.samples import demo_files


@pytest.mark.parametrize("week", [1, 2, 3])
def test_demo_inputs_are_reproducible_across_different_generation_times(monkeypatch, week):
    clock = {"value": datetime(2030, 1, 1)}

    class Clock:
        @staticmethod
        def now(tz=None):
            return clock["value"].replace(tzinfo=tz)

    monkeypatch.setattr(excel, "datetime", SimpleNamespace(datetime=Clock, timezone=timezone))
    first = demo_files(week)
    clock["value"] = datetime(2031, 12, 31, 23, 59, 59)
    second = demo_files(week)
    assert first == second
    for asset in first:
        if asset["name"].endswith(".xlsx"):
            with ZipFile(io.BytesIO(asset["content"])) as archive:
                assert len({member.date_time for member in archive.infolist()}) == 1
                assert b"2030" not in archive.read("docProps/core.xml")
                assert b"2031" not in archive.read("docProps/core.xml")
            workbook = load_workbook(io.BytesIO(asset["content"]))
            assert workbook.active.max_row == (4 if asset["role"] == "master" else 1)
            workbook.close()
