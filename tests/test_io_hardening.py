"""Independent I/O hardening contract tests.

Synthetic data only. These tests exercise the parser/exporter directly with
temporary directories; they do not import the API, read configuration or call a
model provider.
"""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, Side

from wheat_assistant import exporter, files

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "io_hardening"
TEMPLATE = FIXTURES / "template_two_sheets.xlsx"
MULTILINE = FIXTURES / "multiline_orders.csv"

EXPECTED_SHEETS = [
    "Warehouse_Picklist",
    "Order_Detail",
    "Exceptions",
    "Excluded",
    "Duplicates",
    "SKU_Summary",
    "Lineage",
    "Run_Summary",
]

FLOW = {
    "output_columns": [
        {"header": "SKU", "field": "sku"},
        {"header": "Description", "field": "name"},
        {"header": "Pick quantity", "field": "quantity"},
        {"header": "Unit", "field": "unit"},
    ]
}


def make_csv(headers: list[str], rows: list[list[str]]) -> bytes:
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8")


def sample_result(
    *,
    name: str = "Synthetic tea",
    quantity: str = "3",
    note: str = "synthetic note",
) -> dict:
    raw = {"Order": "SYN-1", "Note": note}
    rows = [
        {
            "source_row_id": "synthetic:2",
            "source_id": "channel-a",
            "source_name": "synthetic_orders.csv",
            "sheet": "CSV",
            "row_number": 2,
            "sku": "TEA-001",
            "product_name": name,
            "quantity": quantity,
            "unit": "piece",
            "disposition": "included",
            "reasons": [],
            "events": [],
            "duplicate_of": None,
            "raw": raw,
        },
        {
            "source_row_id": "synthetic:3",
            "source_id": "channel-a",
            "source_name": "synthetic_orders.csv",
            "sheet": "CSV",
            "row_number": 3,
            "sku": None,
            "product_name": "",
            "quantity": None,
            "unit": "",
            "disposition": "exception",
            "reasons": [{"code": "SYNTHETIC_REASON", "message": "synthetic reason"}],
            "events": [],
            "duplicate_of": None,
            "raw": {"Order": "SYN-2", "Note": "synthetic exception"},
        },
    ]
    return {
        "rows": rows,
        "summary": {
            "input": 2,
            "included": 1,
            "excluded": 0,
            "duplicate": 0,
            "exception": 1,
            "totals_match": True,
        },
        "sku_summary": [
            {"sku": "TEA-001", "name": name, "quantity": quantity, "unit": "piece"}
        ],
    }


def test_parser_version_and_valid_multiline_csv_line_spans() -> None:
    assert files.PARSER_VERSION == 2
    table = files.parse_table(
        MULTILINE.read_bytes(),
        "synthetic_multiline.csv",
        "orders",
    )

    assert table["parser_version"] == 2
    assert table["headers"] == ["Order", "Line", "Note"]
    assert table["sheet"] == "CSV"
    records = table["records"]
    assert len(records) == 3
    assert [
        (record["row_number"], record["row_end_number"]) for record in records
    ] == [(2, 3), (4, 4), (5, 5)]
    assert [record["data"]["Note"] for record in records] == [
        "first line\nsecond line",
        'say "hello", now',
        "plain",
    ]
    assert [record["data"]["Order"] for record in records] == [
        "SYN-1",
        "SYN-2",
        "SYN-3",
    ]

    profiled = files.profile(table)
    assert profiled["row_count"] == 3
    assert len(profiled["samples"]) == 3
    assert profiled["samples"][0]["Note"] == "first line\nsecond line"


@pytest.mark.parametrize(
    "content",
    [
        pytest.param(
            b'Order,Line,Note\nSYN-1,1,"unclosed\nSYN-2,2,swallowed\n',
            id="unclosed-quote",
        ),
        pytest.param(
            b'Order,Line,Note\nSYN-1,1,"closed"junk\n',
            id="characters-after-closing-quote",
        ),
        pytest.param(
            b'Order,Line,Note\nSYN-1,1,bad"quote\n',
            id="quote-inside-unquoted-field",
        ),
    ],
)
def test_malformed_csv_quotes_raise_safe_value_error(content: bytes) -> None:
    with pytest.raises(ValueError) as caught:
        files.parse_table(content, "synthetic_malformed.csv", "orders")

    error = caught.value
    message = str(error)
    assert "synthetic_malformed.csv" in message
    assert error.code == "INVALID_CSV"
    assert isinstance(error.line_number, int) and error.line_number >= 1
    assert "/" not in message


def test_csv_cell_limit_is_32767_and_overlong_values_are_rejected() -> None:
    assert files.MAX_CELL_CHARS == 32767

    exact_header = "H" * files.MAX_CELL_CHARS
    accepted = files.parse_table(
        (exact_header + "\n").encode("utf-8"),
        "synthetic_boundary.csv",
        "template",
    )
    assert accepted["headers"] == [exact_header]
    assert accepted["records"] == []

    for length in (files.MAX_CELL_CHARS + 1, 150000):
        content = make_csv(["Order", "Note"], [["SYN-1", "x" * length]])
        with pytest.raises(ValueError) as caught:
            files.parse_table(content, "synthetic_long.csv", "orders")
        error = caught.value
        assert error.code == "CELL_TOO_LONG"
        assert error.line_number == 2
        assert "synthetic_long.csv" in str(error)
        assert "/" not in str(error)


def test_serialized_raw_data_beyond_excel_limit_is_rejected() -> None:
    headers = [f"Field{i}" for i in range(5)]
    row = ["x" * 8000 for _ in headers]
    content = make_csv(headers, [row])

    with pytest.raises(ValueError) as caught:
        files.parse_table(content, "synthetic_wide_raw.csv", "orders")

    error = caught.value
    assert error.code == "RAW_TOO_LONG"
    assert error.line_number == 2
    assert "synthetic_wide_raw.csv" in str(error)


def test_xlsx_cell_limit_and_physical_row_numbers() -> None:
    overlong = Workbook()
    overlong_sheet = overlong.active
    overlong_sheet.title = "Orders"
    overlong_sheet.append(["Order", "Note"])
    overlong_sheet.append(["SYN-1", "x" * files.MAX_CELL_CHARS])
    stream = io.BytesIO()
    overlong.save(stream)
    patched = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(
        patched, "w"
    ) as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b"x" * files.MAX_CELL_CHARS,
                    b"x" * (files.MAX_CELL_CHARS + 1),
                )
            target.writestr(item, data)

    with pytest.raises(ValueError) as caught:
        files.parse_table(patched.getvalue(), "synthetic_xlsx_long.xlsx", "orders")
    assert caught.value.code == "CELL_TOO_LONG"
    assert caught.value.line_number == 2

    workbook = Workbook()
    intro = workbook.active
    intro.title = "Intro"
    intro.append(["Ignore this sheet"])
    orders = workbook.create_sheet("Orders")
    orders.append(["Order", "Note"])
    orders.append([])
    orders.append(["SYN-2", "second physical row"])
    orders.append(["SYN-3", "third physical row"])
    stream = io.BytesIO()
    workbook.save(stream)

    table = files.parse_table(
        stream.getvalue(),
        "synthetic_physical_rows.xlsx",
        "orders",
        sheet_name="Orders",
    )
    assert table["sheet"] == "Orders"
    assert [
        (record["row_number"], record["row_end_number"])
        for record in table["records"]
    ] == [(3, 3), (4, 4)]


def test_export_workbook_legacy_signature_versions_and_data(tmp_path: Path) -> None:
    assert exporter.EXPORTER_VERSION == 2
    result = sample_result(name="=1+1", quantity="1234567890123456789")
    target = tmp_path / "legacy.xlsx"

    returned = exporter.export_workbook(result, FLOW, "synthetic-snapshot", target)
    assert returned == str(target)
    assert target.is_file()

    workbook = load_workbook(target, data_only=False)
    try:
        assert workbook.sheetnames == EXPECTED_SHEETS
        pick = workbook["Warehouse_Picklist"]
        assert pick["A2"].value == "TEA-001"
        assert pick["B2"].value == "=1+1"
        assert pick["B2"].data_type == "s"
        assert pick["C2"].value == "1234567890123456789"
        assert pick["C2"].data_type == "s"

        assert workbook["Order_Detail"].max_row == 2
        assert workbook["Exceptions"].max_row == 2
        assert workbook["Excluded"].max_row == 1
        assert workbook["Duplicates"].max_row == 1
        assert workbook["SKU_Summary"].max_row == 2
        assert workbook["Lineage"].max_row == 3
        assert workbook["Run_Summary"].max_row == 10
    finally:
        workbook.close()


def test_template_header_styles_use_public_attributes(tmp_path: Path) -> None:
    target = tmp_path / "styled.xlsx"
    exporter.export_workbook(
        sample_result(),
        FLOW,
        "synthetic-snapshot",
        target,
        TEMPLATE.read_bytes(),
    )

    workbook = load_workbook(target, data_only=False)
    try:
        pick = workbook["Warehouse_Picklist"]
        for column in ("A1", "B1", "C1", "D1"):
            cell = pick[column]
            assert cell.border.bottom.style == "thick"
            assert cell.border.bottom.color.rgb == "FF112233"
            assert cell.border.top.style == "double"
            assert cell.border.top.color.rgb == "FF445566"
            assert cell.font.name == "Arial"
            assert cell.font.bold is True
            assert cell.font.italic is True
            assert cell.font.color.rgb == "FF123456"
            assert cell.fill.fgColor.rgb == "FFABCDEF"
            assert cell.alignment.horizontal == "center"
            assert cell.alignment.vertical == "center"
            assert cell.alignment.wrap_text is True
            assert cell.number_format == "@"
        assert pick.column_dimensions["A"].width == 31
        assert pick.column_dimensions["D"].width == 34
        assert pick.row_dimensions[1].height == 42
        assert pick["A1"].fill.fgColor.rgb != "343A35"
    finally:
        workbook.close()


def test_template_sheet_name_selects_the_second_sheet(tmp_path: Path) -> None:
    target = tmp_path / "alternate.xlsx"
    exporter.export_workbook(
        sample_result(),
        FLOW,
        "synthetic-snapshot",
        target,
        TEMPLATE.read_bytes(),
        template_sheet_name="Alternate Header",
    )

    workbook = load_workbook(target, data_only=False)
    try:
        pick = workbook["Warehouse_Picklist"]
        cell = pick["A1"]
        assert cell.fill.fgColor.rgb == "FF445566"
        assert cell.border.bottom.style == "dashed"
        assert cell.border.bottom.color.rgb == "FF778899"
        assert cell.border.top.style == "hair"
        assert cell.font.name == "Calibri"
        assert cell.font.italic is False
        assert cell.alignment.horizontal == "right"
        assert cell.number_format == "0.00"
        assert pick.column_dimensions["A"].width == 21
        assert pick.row_dimensions[1].height == 33
    finally:
        workbook.close()



def test_template_header_without_height_still_keeps_custom_style(
    tmp_path: Path,
) -> None:
    template = Workbook()
    source = template.active
    source.title = "Chosen Header"
    source.append(["SKU", "Description", "Pick quantity", "Unit"])
    for cell in source[1]:
        cell.font = Font(name="Courier New", bold=True, color="FF0A0B0C")
        cell.border = Border(bottom=Side(style="thin", color="FF0D0E0F"))
        cell.number_format = "@"
    source.column_dimensions["A"].width = 27
    stream = io.BytesIO()
    template.save(stream)
    template.close()

    target = tmp_path / "no-height.xlsx"
    exporter.export_workbook(
        sample_result(),
        FLOW,
        "synthetic-snapshot",
        target,
        stream.getvalue(),
        template_sheet_name="Chosen Header",
    )

    workbook = load_workbook(target, data_only=False)
    try:
        cell = workbook["Warehouse_Picklist"]["A1"]
        assert cell.font.name == "Courier New"
        assert cell.font.bold is True
        assert cell.font.color.rgb == "FF0A0B0C"
        assert cell.border.bottom.style == "thin"
        assert cell.border.bottom.color.rgb == "FF0D0E0F"
        assert cell.fill.patternType is None
        assert workbook["Warehouse_Picklist"].column_dimensions["A"].width == 27
    finally:
        workbook.close()


def test_missing_template_sheet_name_raises_value_error(tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        exporter.export_workbook(
            sample_result(),
            FLOW,
            "synthetic-snapshot",
            tmp_path / "missing-sheet.xlsx",
            TEMPLATE.read_bytes(),
            template_sheet_name="Synthetic Missing Sheet",
        )


def test_export_stage_callbacks_are_writing_then_verifying(tmp_path: Path) -> None:
    stages: list[str] = []
    target = tmp_path / "stages.xlsx"

    exporter.export_workbook(
        sample_result(),
        FLOW,
        "synthetic-snapshot",
        target,
        on_stage=stages.append,
    )

    assert stages == ["writing", "verifying"]
    assert target.is_file()


def test_failed_verification_preserves_target_and_cleans_temp(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target = tmp_path / "atomic.xlsx"
    original = b"ORIGINAL-SYNTHETIC-TARGET"
    target.write_bytes(original)
    stages: list[str] = []

    def fail_verification(*args, **kwargs):
        raise ValueError("synthetic verification failure")

    monkeypatch.setattr(exporter, "load_workbook", fail_verification)

    with pytest.raises(ValueError, match="verification"):
        exporter.export_workbook(
            sample_result(),
            FLOW,
            "synthetic-snapshot",
            target,
            on_stage=stages.append,
        )

    assert target.read_bytes() == original
    assert stages == ["writing", "verifying"]
    assert sorted(path.name for path in tmp_path.iterdir()) == ["atomic.xlsx"]
