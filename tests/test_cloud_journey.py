"""Real API flow over scoped SQL storage; optionally use Replit PostgreSQL in situ."""
import hashlib
import io
import os
import shutil

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine

from wheat_assistant.cloud import create_cloud_app
from wheat_assistant.cloud_store import CloudDatabase

HOST = "wheat-demo.replit.app"
HEADERS = {"Origin": "https://" + HOST, "X-Wheat-Request": "1"}


def checked(response):
    assert response.status_code == 200, response.text
    return response.json()


def test_three_weeks_restart_download_and_second_browser(tmp_path, monkeypatch):
    monkeypatch.setenv("WHEAT_CLOUD_LIVE_AI", "0")
    monkeypatch.setenv("LLM_API_KEY", "")
    url = os.getenv("WHEAT_TEST_DATABASE_URL")
    engine = None if url else create_engine("sqlite:///" + str(tmp_path / "cloud.sqlite"), connect_args={"check_same_thread": False})
    database = CloudDatabase(url, tmp_path / "cache", engine=engine)
    gateway = create_cloud_app(database=database, hosts=[HOST])
    saved_runs, downloads = [], []
    with TestClient(gateway, base_url="https://" + HOST, headers=HEADERS) as first:
        assert first.get("/api/state").status_code == 401
        checked(first.post("/api/session"))
        assert first.get("/healthz").status_code == 200
        assert first.get("/api/health").json()["mode"] == "cloud"
        source = None
        for week, expected in [(1, (36, 24, 3, 3, 6)), (2, (39, 24, 6, 3, 6)), (3, (24, 24, 0, 0, 0))]:
            batch = checked(first.post("/api/demo/load", json={"week": week}))
            draft = checked(first.post(f"/api/batches/{batch['id']}/plan", json={"mode": "suggested", "source_flow_id": source}))
            flow = draft["flow"]
            flow["aliases"]["Jasmine500"] = "TEA-001"
            flow["unit_aliases"]["carton"] = "box"
            if week >= 2:
                flow["channels"]["channel-b"]["status_actions"]["Partially shipped"] = "exclude"
            report = checked(first.post(f"/api/batches/{batch['id']}/preflight", json=flow))
            approval = checked(first.post(f"/api/batches/{batch['id']}/approve", json={"flow": flow,
                "reviewed": ["fields", "statuses", "products", "units", "duplicates", "output"],
                "expected_input_hash": report["input_hash"], "source_flow_id": source, "draft_id": draft["id"]}))
            source = approval["id"]
            run = checked(first.post("/api/runs", json={"batch_id": batch["id"], "flow_id": source}))
            completed = checked(first.post(f"/api/runs/{run['id']}/execute"))
            assert tuple(completed["summary"][k] for k in ("input", "included", "excluded", "duplicate", "exception")) == expected
            assert first.get(f"/api/runs/{run['id']}/download").status_code == 403
            export = checked(first.post(f"/api/runs/{run['id']}/export-approval", json={
                "result_hash": completed["result_hash"], "acknowledge_exceptions": True}))
            artifact = first.get(export["download_url"])
            assert artifact.status_code == 200
            book = load_workbook(io.BytesIO(artifact.content))
            assert len(book.sheetnames) == 8
            book.close()
            assert checked(first.post(f"/api/runs/{run['id']}/handoff"))["simulation"] is True
            saved_runs.append(run["id"])
            downloads.append(hashlib.sha256(artifact.content).hexdigest())
        original_id = batch["files"][0]["id"]
        original = first.get(f"/api/files/{original_id}/download").content
        cookie = first.cookies
        assert first.post(f"/api/batches/{batch['id']}/plan", json={"mode": "live"}).status_code == 503

    # Simulate complete loss of the deployment's ephemeral disk, not just a browser refresh.
    shutil.rmtree(database.cache_root)
    restarted = create_cloud_app(database=database, hosts=[HOST])
    with TestClient(restarted, base_url="https://" + HOST, headers=HEADERS, cookies=cookie) as resumed:
        assert len(checked(resumed.get("/api/state"))["runs"]) == 3
        for identifier, checksum in zip(saved_runs, downloads):
            response = resumed.get(f"/api/runs/{identifier}/download")
            assert response.status_code == 200
            assert hashlib.sha256(response.content).hexdigest() == checksum
        assert resumed.get(f"/api/files/{original_id}/download").content == original
        with TestClient(restarted, base_url="https://" + HOST, headers=HEADERS) as other:
            checked(other.post("/api/session"))
            assert checked(other.get("/api/state"))["runs"] == []
            for path in [f"/api/runs/{saved_runs[0]}", f"/api/runs/{saved_runs[0]}/download", f"/api/files/{original_id}/download"]:
                assert other.get(path).status_code == 404


def test_desktop_still_rejects_remote_hosts(tmp_path):
    from wheat_assistant.api import create_app
    with TestClient(create_app(tmp_path), base_url="https://" + HOST) as client:
        assert client.get("/api/state").status_code == 403
