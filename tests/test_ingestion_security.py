"""Independent ingestion, approval and export security tests.

These tests run against the real FastAPI app with TestClient and an isolated
``tmp_path`` store. All inputs are synthetic. No external model is contacted:
LLM environment variables are cleared, dotenv loading is disabled before the API
module import, and ``propose_live`` is replaced with a hard failure.
"""

from __future__ import annotations

import copy
import csv
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

REVIEW_GROUPS = ["fields", "statuses", "products", "units", "duplicates", "output"]
ALIASES = {"Jasmine500": "TEA-001"}
UNIT_ALIASES = {"carton": "box"}
LLM_ENV = ("LLM_BASE_URL", "LLM_MODEL", "LLM_API_KEY")


@pytest.fixture()
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    for name in LLM_ENV:
        monkeypatch.setenv(name, "")

    import dotenv

    monkeypatch.setattr(dotenv, "load_dotenv", lambda *args, **kwargs: False)

    from wheat_assistant import api

    def forbidden_live(*args, **kwargs):
        raise AssertionError("Live model calls are forbidden in these tests")

    monkeypatch.setattr(api, "propose_live", forbidden_live)
    app = api.create_app(data_dir=tmp_path)
    with TestClient(app) as test_client:
        yield test_client


def csv_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    target = io.StringIO()
    writer = csv.writer(target)
    writer.writerow(headers)
    writer.writerows(rows)
    return target.getvalue().encode("utf-8-sig")


def xlsx_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def upload(
    client: TestClient,
    filename: str,
    content: bytes,
    *,
    role: str,
    channel_id: str = "",
):
    return client.post(
        "/api/files",
        files={"file": (filename, content, "application/octet-stream")},
        data={"role": role, "channel_id": channel_id},
    )


def load_week(client: TestClient, week: int) -> dict:
    response = client.post("/api/demo/load", json={"week": week})
    assert response.status_code == 200, response.text
    return response.json()


def suggested_flow(
    client: TestClient,
    batch_id: str,
    *,
    apply_aliases: bool = True,
) -> dict:
    response = client.post(f"/api/batches/{batch_id}/plan", json={"mode": "suggested"})
    assert response.status_code == 200, response.text
    plan = response.json()
    assert plan["origin"] == "suggested_setup"
    assert plan["model"] is None
    flow = copy.deepcopy(plan["flow"])
    if apply_aliases:
        flow["aliases"] = dict(ALIASES)
        flow["unit_aliases"] = dict(UNIT_ALIASES)
    return flow


def preflight(client: TestClient, batch_id: str, flow: dict) -> dict:
    response = client.post(f"/api/batches/{batch_id}/preflight", json=flow)
    assert response.status_code == 200, response.text
    return response.json()


def approve(
    client: TestClient,
    batch_id: str,
    flow: dict,
    input_hash: str,
    reviewed: list[str] | None = None,
) -> dict:
    response = client.post(
        f"/api/batches/{batch_id}/approve",
        json={
            "flow": flow,
            "reviewed": list(REVIEW_GROUPS if reviewed is None else reviewed),
            "expected_input_hash": input_hash,
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def execute(client: TestClient, batch_id: str, flow_id: str) -> dict:
    response = client.post(
        "/api/runs",
        json={"batch_id": batch_id, "flow_id": flow_id},
    )
    assert response.status_code == 200, response.text
    run = response.json()
    assert run["status"] == "ready"

    response = client.post(f"/api/runs/{run['id']}/execute")
    assert response.status_code == 200, response.text
    completed = response.json()
    assert completed["status"] == "completed"
    return completed


def completed_week1_run(client: TestClient) -> dict:
    batch = load_week(client, 1)
    flow = suggested_flow(client, batch["id"])
    report = preflight(client, batch["id"], flow)
    assert report["errors"] == []
    saved = approve(client, batch["id"], flow, report["input_hash"])
    run = execute(client, batch["id"], saved["id"])
    assert run["summary"] == {
        "input": 36,
        "included": 24,
        "excluded": 3,
        "duplicate": 3,
        "exception": 6,
        "totals_match": True,
    }
    return run


def test_cross_origin_writes_are_rejected(client: TestClient) -> None:
    blocked = client.post(
        "/api/demo/load",
        json={"week": 1},
        headers={"Origin": "https://evil.example"},
    )
    assert blocked.status_code == 403, blocked.text
    assert "Cross-origin writes" in blocked.json()["detail"]
    assert client.get("/api/state").json()["batches"] == []

    allowed = client.post("/api/demo/load", json={"week": 1})
    assert allowed.status_code == 200, allowed.text
    assert len(allowed.json()["files"]) == 5
    assert len(client.get("/api/state").json()["batches"]) == 1


def test_csv_with_duplicate_headers_is_rejected(client: TestClient) -> None:
    content = csv_bytes(
        ["Order ID", "Line", "Order ID"],
        [["SYN-SEC-1", "1", "duplicate"]],
    )
    response = upload(
        client,
        "duplicate_headers.csv",
        content,
        role="orders",
        channel_id="channel-a",
    )

    assert response.status_code == 422, response.text
    assert "Headers" in response.json()["detail"]
    assert client.get("/api/activity").json() == []


def test_non_utf8_csv_is_rejected(client: TestClient) -> None:
    content = b"Order ID,Line\n\xff\xfe,1\n"
    response = upload(
        client,
        "non_utf8.csv",
        content,
        role="orders",
        channel_id="channel-a",
    )

    assert response.status_code == 422, response.text
    assert "UTF-8" in response.json()["detail"]
    assert client.get("/api/activity").json() == []


def test_corrupt_xlsx_is_rejected(client: TestClient) -> None:
    response = upload(
        client,
        "corrupt.xlsx",
        b"this is not a zip archive",
        role="orders",
        channel_id="channel-a",
    )

    assert response.status_code == 422, response.text
    assert "Workbook cannot be read" in response.json()["detail"]
    assert client.get("/api/activity").json() == []


def test_formula_in_product_master_is_rejected(client: TestClient) -> None:
    content = xlsx_bytes(
        ["SKU", "Name", "Base unit", "Box size", "Pack size"],
        [["TEA-001", '=UPPER("tea")', "piece", "12", "6"]],
    )
    response = upload(
        client,
        "master_with_formula.xlsx",
        content,
        role="master",
    )

    assert response.status_code == 422, response.text
    assert "formula" in response.json()["detail"].lower()
    assert client.get("/api/activity").json() == []


def test_approve_rejects_stale_input_hash(client: TestClient) -> None:
    week1 = load_week(client, 1)
    week1_flow = suggested_flow(client, week1["id"])
    week1_report = preflight(client, week1["id"], week1_flow)

    week2 = load_week(client, 2)
    week2_flow = suggested_flow(client, week2["id"])
    week2_report = preflight(client, week2["id"], week2_flow)
    assert week1_report["input_hash"] != week2_report["input_hash"]

    response = client.post(
        f"/api/batches/{week2['id']}/approve",
        json={
            "flow": week2_flow,
            "reviewed": REVIEW_GROUPS,
            "expected_input_hash": week1_report["input_hash"],
        },
    )

    assert response.status_code == 422, response.text
    assert "Inputs changed" in response.json()["detail"]
    assert client.get("/api/state").json()["flows"] == []


def test_approve_rejects_each_missing_review_group(client: TestClient) -> None:
    batch = load_week(client, 1)
    flow = suggested_flow(client, batch["id"])
    report = preflight(client, batch["id"], flow)

    for missing in REVIEW_GROUPS:
        reviewed = [group for group in REVIEW_GROUPS if group != missing]
        response = client.post(
            f"/api/batches/{batch['id']}/approve",
            json={
                "flow": flow,
                "reviewed": reviewed,
                "expected_input_hash": report["input_hash"],
            },
        )
        assert response.status_code == 422, (missing, response.text)
        assert "Review each concrete rule group" in response.json()["detail"]

    assert client.get("/api/state").json()["flows"] == []


def test_download_is_blocked_until_export_approval_and_exceptions_need_ack(
    client: TestClient,
) -> None:
    run = completed_week1_run(client)
    assert run["summary"]["exception"] == 6

    blocked = client.get(f"/api/runs/{run['id']}/download")
    assert blocked.status_code == 403, blocked.text
    assert "Approve the reviewed result" in blocked.json()["detail"]

    no_ack = client.post(
        f"/api/runs/{run['id']}/export-approval",
        json={"result_hash": run["result_hash"], "acknowledge_exceptions": False},
    )
    assert no_ack.status_code == 422, no_ack.text
    assert "Acknowledge" in no_ack.json()["detail"]
    assert client.get(f"/api/runs/{run['id']}/download").status_code == 403

    approved = client.post(
        f"/api/runs/{run['id']}/export-approval",
        json={"result_hash": run["result_hash"], "acknowledge_exceptions": True},
    )
    assert approved.status_code == 200, approved.text
    artifact = client.get(approved.json()["download_url"])
    assert artifact.status_code == 200, artifact.text
    assert artifact.content.startswith(b"PK")
    assert len(artifact.content) > 1000


def test_handoff_is_idempotent_and_requires_export_approval(
    client: TestClient,
) -> None:
    run = completed_week1_run(client)

    premature = client.post(f"/api/runs/{run['id']}/handoff")
    assert premature.status_code == 422, premature.text
    assert "Approve the output" in premature.json()["detail"]

    approved = client.post(
        f"/api/runs/{run['id']}/export-approval",
        json={"result_hash": run["result_hash"], "acknowledge_exceptions": True},
    )
    assert approved.status_code == 200, approved.text

    first = client.post(f"/api/runs/{run['id']}/handoff")
    second = client.post(f"/api/runs/{run['id']}/handoff")
    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    task = first.json()
    assert second.json() == task
    assert task["run_id"] == run["id"]
    assert task["snapshot"] == run["result_hash"]
    assert task["simulation"] is True
    assert task["status"] == "ready_for_review"


def test_export_keeps_leading_equals_text_as_a_string(client: TestClient) -> None:
    orders = csv_bytes(
        ["Order ID", "Line", "Item Code", "Qty", "Unit", "Order Status", "Date"],
        [["SYN-FORMULA-1", "1", "TEA-001", "1", "piece", "Confirmed", "2026-10-06"]],
    )
    master = csv_bytes(
        ["SKU", "Name", "Base unit", "Box size", "Pack size"],
        [["TEA-001", "=1+1", "piece", "12", "6"]],
    )

    order_asset = upload(
        client,
        "formula_orders.csv",
        orders,
        role="orders",
        channel_id="channel-a",
    )
    assert order_asset.status_code == 200, order_asset.text
    master_asset = upload(
        client,
        "formula_master.csv",
        master,
        role="master",
    )
    assert master_asset.status_code == 200, master_asset.text

    created = client.post(
        "/api/batches",
        json={
            "file_ids": [order_asset.json()["id"], master_asset.json()["id"]],
            "label": "Formula text export",
        },
    )
    assert created.status_code == 200, created.text
    batch = created.json()

    flow = suggested_flow(client, batch["id"], apply_aliases=False)
    report = preflight(client, batch["id"], flow)
    assert report["row_count"] == 1
    assert report["errors"] == []
    assert report["unresolved_products"] == {}
    assert report["unresolved_units"] == {}

    saved = approve(client, batch["id"], flow, report["input_hash"])
    run = execute(client, batch["id"], saved["id"])
    assert run["summary"]["exception"] == 0
    assert run["summary"]["included"] == 1

    approved = client.post(
        f"/api/runs/{run['id']}/export-approval",
        json={"result_hash": run["result_hash"], "acknowledge_exceptions": False},
    )
    assert approved.status_code == 200, approved.text
    artifact = client.get(approved.json()["download_url"])
    assert artifact.status_code == 200, artifact.text

    workbook = load_workbook(io.BytesIO(artifact.content), data_only=False)
    try:
        pick = workbook["Warehouse_Picklist"]
        assert pick["A2"].value == "TEA-001"
        assert pick["B2"].value == "=1+1"
        assert pick["B2"].data_type == "s"
        assert pick["B2"].number_format == "@"
        assert pick["C2"].value == 1

        summary = workbook["SKU_Summary"]
        assert summary["B2"].value == "=1+1"
        assert summary["B2"].data_type == "s"
        assert summary["B2"].number_format == "@"

        formula_cells = [
            (sheet.title, cell.coordinate, cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        assert formula_cells == []
    finally:
        workbook.close()
