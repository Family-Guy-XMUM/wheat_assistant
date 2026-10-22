"""Run a real HTTP journey against source or an independently extracted native ZIP."""
import argparse
import hashlib
import io
import json
import os
import platform
import socket
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

def verify(bundle=None):
    with tempfile.TemporaryDirectory(prefix="maizi release 合成 ") as temporary:
        folder = Path(temporary)
        if bundle:
            with zipfile.ZipFile(bundle) as archive:
                if any(Path(name).is_absolute() or ".." in Path(name).parts for name in archive.namelist()):
                    raise ValueError("Unsafe bundle member")
                archive.extractall(folder / "native")
            binary = folder / "native" / "MaiZi" / ("MaiZi.exe" if os.name == "nt" else "MaiZi")
            if os.name != "nt":
                binary.chmod(0o755)
            command = [str(binary)]
            environment = {key: value for key, value in os.environ.items()
                           if key in {"SystemRoot", "WINDIR", "TEMP", "TMP", "TMPDIR", "LANG", "LC_ALL"}}
            environment["PATH"] = os.path.join(os.environ.get("SystemRoot", "C:\\Windows"), "System32") if os.name == "nt" else "/usr/bin:/bin"
        else:
            command = [sys.executable, "-m", "maizi.launcher"]
            environment = os.environ.copy()
        environment.update({"LLM_API_KEY": "", "LLM_BASE_URL": "", "LLM_MODEL": "", "PYTHONUTF8": "1"})
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            port = probe.getsockname()[1]
        command += ["--no-browser", "--port", str(port), "--data-dir", str(folder / "data")]
        process = subprocess.Popen(command, cwd=folder, env=environment, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        started = time.monotonic()
        try:
            with httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=60) as client:
                for _ in range(400):
                    if process.poll() is not None:
                        raise RuntimeError("App exited before its health check")
                    try:
                        health = client.get("/api/health")
                        if health.status_code == 200:
                            break
                    except httpx.HTTPError:
                        pass
                    time.sleep(.15)
                else:
                    raise RuntimeError("App startup deadline exceeded")
                assert health.json()["ai_configured"] is False
                duplicate=subprocess.run(command,cwd=folder,env=environment,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=30)
                assert duplicate.returncode!=0, "Duplicate desktop instance was not blocked"
                assert client.get("/api/health").status_code==200
                def post(path, data=None):
                    response = client.post(path, json=data) if data is not None else client.post(path)
                    if response.status_code != 200:
                        raise RuntimeError("Verification request failed: " + path.split("?")[0])
                    return response.json()
                source = None
                runs = []
                for week, expected in [(1,(36,24,3,3,6)),(2,(39,24,6,3,6)),(3,(24,24,0,0,0))]:
                    batch = post("/api/demo/load", {"week":week})
                    draft = post("/api/batches/"+batch["id"]+"/plan", {"mode":"suggested","source_flow_id":source})
                    spec = draft["flow"]
                    spec["aliases"]["Jasmine500"] = "TEA-001"
                    spec["unit_aliases"]["carton"] = "box"
                    if week == 2:
                        pre = post("/api/batches/"+batch["id"]+"/preflight", spec)
                        assert len(pre["status_guard"]["affected_row_ids"]) == 3
                        spec["channels"]["channel-b"]["status_actions"]["Partially shipped"] = "exclude"
                    pre = post("/api/batches/"+batch["id"]+"/preflight", spec)
                    saved = post("/api/batches/"+batch["id"]+"/approve", {"flow":spec,"source_flow_id":source,
                        "expected_input_hash":pre["input_hash"],"reviewed":["fields","statuses","products","units","duplicates","output"]})
                    source = saved["id"]
                    run = post("/api/runs", {"batch_id":batch["id"],"flow_id":source})
                    done = post("/api/runs/"+run["id"]+"/execute")
                    assert tuple(done["summary"][key] for key in ["input","included","excluded","duplicate","exception"]) == expected
                    approved = post("/api/runs/"+run["id"]+"/export-approval",
                                    {"result_hash":done["result_hash"],"acknowledge_exceptions":True})
                    content = client.get(approved["download_url"])
                    assert content.status_code == 200
                    book = load_workbook(io.BytesIO(content.content),read_only=True)
                    assert len(book.sheetnames) == 8 and book["Lineage"].max_row == expected[0]+1
                    quantities = {row[0]:row[2] for row in book["SKU_Summary"].iter_rows(min_row=2,values_only=True)}
                    assert quantities == {"RICE-003":84,"SOAP-002":33,"TEA-001":21}
                    book.close()
                    task = post("/api/runs/"+run["id"]+"/handoff")
                    assert task["simulation"] is True
                    assert post("/api/runs/"+run["id"]+"/handoff")["id"] == task["id"]
                    runs.append({"week":week,"summary":done["summary"],"rule_version":saved["rule_version"]})
                source_files=sorted(list((ROOT/"backend"/"maizi").rglob("*.py"))+[p for p in (ROOT/"frontend"/"dist").rglob("*") if p.is_file()])
                manifest={p.relative_to(ROOT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in source_files}
                result = {"passed":True,"mode":"native" if bundle else "source","os":platform.system(),
                          "source_sha256":hashlib.sha256(json.dumps(manifest,sort_keys=True).encode()).hexdigest(),
                          "bundle_sha256":hashlib.sha256(Path(bundle).read_bytes()).hexdigest() if bundle else None,
                          "architecture":platform.machine(),"ai_calls":0,"weeks":runs,
                          "duplicate_instance_blocked":True,
                          "seconds":round(time.monotonic()-started,3)}
        finally:
            process.terminate()
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill();process.wait(timeout=5)
        with socket.socket() as stopped:
            assert stopped.connect_ex(("127.0.0.1", port)) != 0, "Port still responds after stopping the app"
        result["stop_and_port_release"] = True
        restarted=subprocess.Popen(command,cwd=folder,env=environment,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        try:
            with httpx.Client(base_url=f"http://127.0.0.1:{port}",timeout=30) as client:
                for _ in range(400):
                    try:
                        if client.get("/api/health").status_code==200:break
                    except httpx.HTTPError:pass
                    if restarted.poll() is not None:raise RuntimeError("Restart exited early")
                    time.sleep(.15)
                else:raise RuntimeError("Restart deadline exceeded")
                historical=client.get("/api/history/runs").json()
                assert historical["total"]==3 and all(x["status"]=="completed" for x in historical["rows"])
                assert client.get(approved["download_url"]).content==content.content
                result["restart_history_and_download"]=True
        finally:
            restarted.terminate()
            try:restarted.wait(timeout=10)
            except subprocess.TimeoutExpired:restarted.kill();restarted.wait(timeout=5)
        return result

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--bundle", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        result = verify(args.bundle)
    except Exception as exc:
        safe_message=str(exc) if isinstance(exc,RuntimeError) and str(exc).startswith(("App ","Restart ","Verification request")) else "Release verification failed; inspect the local environment and rerun."
        result = {"passed":False,"error_type":type(exc).__name__,"message":safe_message}
    if args.output:
        args.output.parent.mkdir(parents=True,exist_ok=True)
        args.output.write_text(json.dumps(result,indent=2),encoding="utf-8")
    print(json.dumps(result))
    raise SystemExit(0 if result["passed"] else 1)
