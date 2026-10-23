"""Three uncached full-size synthetic API runs, including verified resource cleanup."""
import argparse
import csv
import io
import json
import os
import statistics
import tempfile
import threading
import time
import traceback
from pathlib import Path
from unittest.mock import patch

import psutil
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from wheat_assistant.release_evidence import attempt_id, candidate_info, report_path, write_report

for key in ("LLM_API_KEY", "LLM_BASE_URL", "LLM_MODEL"):
    os.environ[key] = ""
with patch("dotenv.load_dotenv", return_value=False):
    from wheat_assistant.api import create_app

ROOT = Path(__file__).resolve().parents[1]
GROUPS = ["fields", "statuses", "products", "units", "duplicates", "output"]
SHEETS = ["Warehouse_Picklist", "Order_Detail", "Exceptions", "Excluded", "Duplicates", "SKU_Summary", "Lineage", "Run_Summary"]


def csv_file(headers, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def trial():
    item = {"status":"running", "phase":"initializing", "cleanup_passed":False, "complete":False,
            "prepare_seconds":None, "execute_seconds":None, "export_seconds":None}
    peak = [0]
    memory_available = [True]
    stopped = threading.Event()
    process = psutil.Process()
    def sample():
        while not stopped.is_set():
            try:
                peak[0] = max(peak[0], sum(p.memory_info().rss for p in [process] + process.children(recursive=True) if p.is_running()))
            except psutil.NoSuchProcess:
                pass
            except (psutil.AccessDenied, PermissionError):
                memory_available[0] = False
            stopped.wait(.025)
    sampler = threading.Thread(target=sample, daemon=True)
    sampler.start()
    try:
        with tempfile.TemporaryDirectory(prefix="wheat_assistant-benchmark-") as folder:
            item["phase"] = "opening_app"
            with TestClient(create_app(Path(folder))) as client:
                def post(url, data):
                    response = client.post(url, json=data)
                    assert response.status_code == 200, response.text
                    return response.json()
                def upload(name, data, role):
                    response = client.post("/api/files", files={"file":(name,data)}, data={"role":role,"channel_id":"bulk" if role=="orders" else ""})
                    assert response.status_code == 200, response.text
                    return response.json()["id"]
                item["phase"] = "preparing_inputs"
                start = time.perf_counter()
                file_ids = []
                for part in range(5):
                    rows = [[f"O{i}","1",f"S{i%5000:04}","2","piece","Confirmed"] for i in range(part*2000,(part+1)*2000)]
                    file_ids.append(upload(f"synthetic-orders-{part}.csv", csv_file(["Order ID","Line","SKU","Qty","Unit","Status"],rows), "orders"))
                master = [[f"S{i:04}",f"Synthetic product {i}","piece","12","6"] for i in range(5000)]
                file_ids.append(upload("synthetic-master.csv", csv_file(["SKU","Name","Base unit","Box size","Pack size"],master), "master"))
                batch = post("/api/batches", {"file_ids":file_ids,"label":"Synthetic scale acceptance"})
                flow = post(f"/api/batches/{batch['id']}/plan", {"mode":"suggested"})["flow"]
                preflight = post(f"/api/batches/{batch['id']}/preflight", flow)
                saved = post(f"/api/batches/{batch['id']}/approve", {"flow":flow,"expected_input_hash":preflight["input_hash"],"reviewed":GROUPS})
                run = post("/api/runs", {"batch_id":batch["id"],"flow_id":saved["id"]})
                prepared = time.perf_counter()
                item.update({"phase":"executing", "prepare_seconds":round(prepared-start,3)})
                done = post(f"/api/runs/{run['id']}/execute", {})
                executed = time.perf_counter()
                item["execute_seconds"] = round(executed-prepared,3)
                assert done["summary"] == {"input":10000,"included":10000,"excluded":0,"duplicate":0,"exception":0,"totals_match":True}
                item["phase"] = "exporting"
                export = post(f"/api/runs/{run['id']}/export-approval", {"result_hash":done["result_hash"],"acknowledge_exceptions":True})
                exported = time.perf_counter()
                item.update({"export_seconds":round(exported-executed,3), "phase":"validating_workbook"})
                response = client.get(export["download_url"])
                assert response.status_code == 200
                book = load_workbook(io.BytesIO(response.content), read_only=True, data_only=False)
                try:
                    assert book.sheetnames == SHEETS
                    counts = {sheet.title:sheet.max_row-1 for sheet in book}
                    assert counts == {"Warehouse_Picklist":5000,"Order_Detail":10000,"Exceptions":0,"Excluded":0,"Duplicates":0,"SKU_Summary":5000,"Lineage":10000,"Run_Summary":9}
                    totals = {r[0]:str(r[2]) for r in book["SKU_Summary"].iter_rows(min_row=2,values_only=True)}
                    assert totals == {f"S{i:04}":"4" for i in range(5000)}
                    lineage = list(book["Lineage"].iter_rows(min_row=2,values_only=True))
                    assert len({r[0] for r in lineage}) == 10000
                    assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)
                finally:
                    book.close()
                item["phase"] = "browsing_all_results"
                all_skus = []
                for offset in range(0,5000,200):
                    page = client.get(f"/api/runs/{run['id']}/sku-summary", params={"offset":offset,"limit":200}).json()
                    assert page["total"] == 5000
                    all_skus.extend(x["sku"] for x in page["rows"])
                assert len(set(all_skus)) == 5000
                last = client.get(f"/api/runs/{run['id']}/rows", params={"offset":9900,"limit":100}).json()
                assert last["total"] == 10000 and len(last["rows"]) == 100
                item.update({"sheet_rows":counts,"unique_lineage":10000,"browsed_skus":len(set(all_skus)),"workbook_bytes":len(response.content),"phase":"closing_app"})
            item["phase"] = "temporary_directory_cleanup"
        item.update({"status":"completed", "phase":"complete", "complete":True, "cleanup_passed":True})
    except Exception as exc:
        private = ROOT / ".runtime" / "benchmark-private" / (attempt_id()+".log")
        private.parent.mkdir(parents=True, exist_ok=True)
        private.write_text(traceback.format_exc(), encoding="utf-8")
        item.update({"status":"failed", "error_type":type(exc).__name__, "os_error":getattr(exc,"errno",None),
                     "winerror":getattr(exc,"winerror",None), "message":"Synthetic benchmark failed at the recorded phase; no incomplete trial counts toward acceptance."})
    finally:
        stopped.set()
        sampler.join(timeout=2)
        item.update({"memory_sampling_available":memory_available[0] and peak[0]>0,
                     "peak_parent_and_children_rss_mib":round(peak[0]/1024**2,1) if peak[0] else None})
    return item


def benchmark():
    result = {**candidate_info(ROOT), "passed":False, "orders":10000, "skus":5000, "files":5,
              "cpu_count":os.cpu_count(),
              "runs":[], "complete_trials":0, "export_median_seconds":None, "target_seconds":10,
              "memory_sampling":"25 ms, parent plus live child RSS; sums shared pages, not an exact allocation measure"}
    failed = False
    for index in range(3):
        item = {"status":"not_run", "complete":False, "reason":"A prior trial failed"} if failed else trial()
        result["runs"].append({"trial":index+1, **item})
        print(json.dumps(result["runs"][-1]), flush=True)
        failed = failed or not item.get("complete", False)
    completed = [x for x in result["runs"] if x.get("complete") and x.get("cleanup_passed")]
    result["complete_trials"] = len(completed)
    if len(completed) == 3:
        result["export_median_seconds"] = statistics.median(x["export_seconds"] for x in completed)
        result["passed"] = result["export_median_seconds"] <= 10 and all(x["memory_sampling_available"] for x in completed)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    target = report_path(ROOT, "performance.json", args.output)
    result = benchmark()
    write_report(target, result)
    print(json.dumps({"passed":result["passed"],"complete_trials":result["complete_trials"],"export_median_seconds":result["export_median_seconds"]}))
    print("Report:", target.relative_to(ROOT) if target.is_relative_to(ROOT) else target.name)
    raise SystemExit(0 if result["passed"] else 1)
